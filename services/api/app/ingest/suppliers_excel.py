# -*- coding: utf-8 -*-
"""قارئ ملف مدد مديونية الموردين.

Reads `مدة مديونية الموردين.xlsx`. Layout (verified on the real file):
  row 1  company name
  row 2  report title
  row 3  headers: تبويب | المدة بالشهر | … | اسم الحساب | … | رقم الحساب
  row 4+ data

Column A = project (تبويب), B = payment term, D = account name, I = account number.
The account number is the key — supplier names repeat across projects
(انجاز الرواد appears in 5 projects, سماء البناء in 4), so names are never used to match.

# م-٢٠ — تحقّق ترويسة وبادئة حساب (لا تثق بفهارس أعمدة ثابتة)

كان هذا القارئ الوحيد في المجموعة (بعكس budget_xlsx وdebts_report_xls وreceivables_excel
وcontractors_balance_xls) الذي يقرأ `COL_NAME`/`COL_ACCOUNT` كفهارس ثابتة بلا أي تحقّق —
ملف «قائمة موظفين» (اسم في العمود D، رقم وظيفي في I تصادف نفس الموضعين) مرّ بنجاح
`issues: []` وأنتج موردين وهميين. الإصلاح طبقتان، كلٌّ منهما وحدها لا تكفي:
  1. **تحقّق ترويسة**: فهارس الأعمدة تُشتقّ من نص الصف الثالث («اسم الحساب»،
     «رقم الحساب») بدل الثقة بموضعها — نفس نمط debts_report_xls._header_indices.
  2. **تحقّق بادئة الحساب**: القاعدة المطلقة في هذا النظام (٢١١ مورد) — أي حساب
     ببادئة أخرى (كرقم وظيفي بالصدفة) يُرفض صفّه صراحة بدل أن يُقبل مورداً وهمياً.
"""
from __future__ import annotations

from typing import Dict, Optional

from app.domain.payables import Supplier, parse_term
from app.ingest.friendly_errors import check_basic_file, describe_excel_open_error

HEADER_ROW = 3          # 1-based؛ يحمل نصوص الترويسة
FIRST_DATA_ROW = 4      # 1-based; rows 1–3 are titles and headers

#: بادئة حساب المورد المطلقة في هذا النظام — انظر «قاعدة البادئة» في CLAUDE.md.
SUPPLIER_ACCOUNT_PREFIX = '211'

#: أسماء الأعمدة كما تظهر حرفياً في الصف الثالث من الملف الحقيقي.
_LABEL_PROJECT = 'تبويب'
_LABEL_TERM = 'المدة بالشهر'
_LABEL_NAME = 'اسم الحساب'
_LABEL_ACCOUNT = 'رقم الحساب'


class SuppliersParseError(Exception):
    pass


def _header_indices(ws) -> Dict[str, Optional[int]]:
    """يشتقّ فهارس الأعمدة من نص ترويسة الصف الثالث بدل الثقة بمواضعها الثابتة."""
    row = next(ws.iter_rows(min_row=HEADER_ROW, max_row=HEADER_ROW, values_only=True), ())

    def find(label: str) -> Optional[int]:
        for i, v in enumerate(row):
            if v is not None and str(v).strip() == label:
                return i
        return None

    return dict(project=find(_LABEL_PROJECT), term=find(_LABEL_TERM),
               name=find(_LABEL_NAME), account=find(_LABEL_ACCOUNT))


def parse(path: str) -> dict:
    """Return {suppliers, issues}."""
    try:
        import openpyxl
    except ImportError as e:      # pragma: no cover
        raise SuppliersParseError('openpyxl is required to read the suppliers file') from e

    check_basic_file(path, 'ملف مدد الموردين', SuppliersParseError)
    try:
        wb = openpyxl.load_workbook(path, data_only=True)
    except Exception as e:
        raise SuppliersParseError(str(describe_excel_open_error(e, path, 'xlsx'))) from e
    ws = wb.active

    idx = _header_indices(ws)
    col_name, col_account = idx['name'], idx['account']
    if col_name is None or col_account is None:
        raise SuppliersParseError(
            'تعذّر التعرّف على ترويسة الملف — الأعمدة المتوقعة («اسم الحساب» و«رقم '
            'الحساب» في الصف الثالث) غير موجودة. تأكد أن هذا فعلاً ملف مدد مديونية '
            'الموردين وليس ملفاً آخر.')
    col_project, col_term = idx['project'], idx['term']

    suppliers: list[Supplier] = []
    issues: list[dict] = []
    seen: dict[str, str] = {}
    wrong_prefix_count = 0

    for r, row in enumerate(ws.iter_rows(min_row=FIRST_DATA_ROW, values_only=True),
                            start=FIRST_DATA_ROW):
        name = row[col_name] if len(row) > col_name else None
        if not name or not str(name).strip():
            continue

        account = str(row[col_account]).strip() if len(row) > col_account and row[col_account] else ''
        if not account:
            issues.append(dict(severity='error', row=r,
                               message=f'مورد بلا رقم حساب: {str(name).strip()} — تم تجاهله'))
            continue

        # قاعدة البادئة (٢١١ مورد) مطلقة — حساب ببادئة أخرى (رقم وظيفي مثلاً) ليس
        # مورداً مهما بدا شكله صحيحاً. انظر تعليق م-٢٠ أعلى الملف.
        if not account.startswith(SUPPLIER_ACCOUNT_PREFIX):
            wrong_prefix_count += 1
            issues.append(dict(severity='error', row=r, kind='wrong_prefix',
                               message=(f'الحساب {account} ({str(name).strip()}) لا يبدأ '
                                       f'بالبادئة {SUPPLIER_ACCOUNT_PREFIX} (بادئة الموردين) '
                                       '— تم تجاهله')))
            continue

        if account in seen:
            issues.append(dict(severity='error', row=r,
                               message=f'رقم الحساب {account} مكرر ({seen[account]}) — تم تجاهل الصف'))
            continue
        seen[account] = str(name).strip()

        term_raw = row[col_term] if col_term is not None and len(row) > col_term else None
        term = parse_term(str(term_raw) if term_raw is not None else None)
        if term.is_claim:
            issues.append(dict(severity='info', row=r,
                               message=f'{str(name).strip()}: مدة «{term.raw}» — يحتاج تاريخ استحقاق يدوي'))

        project_raw = row[col_project] if col_project is not None and len(row) > col_project else None
        suppliers.append(Supplier(
            account=account,
            name=str(name).strip(),
            project=str(project_raw).strip() if project_raw else '',
            term=term,
        ))

    if not suppliers:
        if wrong_prefix_count:
            raise SuppliersParseError(
                f'لم يُقرأ أي مورد — كل الحسابات ذات الاسم في الملف ({wrong_prefix_count}) '
                f'لا تبدأ بالبادئة {SUPPLIER_ACCOUNT_PREFIX} (بادئة الموردين). تأكد أن هذا '
                'فعلاً ملف مدد مديونية الموردين وليس ملفاً آخر (مثل قائمة موظفين).')
        raise SuppliersParseError('لم يُقرأ أي مورد — تأكد من صيغة الملف')

    return dict(suppliers=suppliers, issues=issues)
