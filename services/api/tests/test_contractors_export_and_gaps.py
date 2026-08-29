# -*- coding: utf-8 -*-
"""اختبارات: ورقة تحليل تصدير المقاولين · تدهور الشعار بلا Pillow · تغطية المقاولين ·
فجوة «بلا كشف حركات» في totals/overview · جمع Decimal في receivables_service.
"""
import importlib
import io

import pytest
from openpyxl import load_workbook


@pytest.fixture()
def env(tmp_path, monkeypatch):
    monkeypatch.setenv('EGCO_DATA_DIR', str(tmp_path / 'data'))

    import app.core.config as config_mod
    importlib.reload(config_mod)
    import app.db.session as session_mod
    importlib.reload(session_mod)
    import app.db.models as models_mod
    importlib.reload(models_mod)
    import app.services.contractors_service as contractors_service_mod
    importlib.reload(contractors_service_mod)
    import app.services.export_service as export_service_mod
    importlib.reload(export_service_mod)
    import app.services.coverage_service as coverage_service_mod
    importlib.reload(coverage_service_mod)
    import app.services.overview_service as overview_service_mod
    importlib.reload(overview_service_mod)
    import app.api.routes.contractors as contractors_route
    importlib.reload(contractors_route)

    session_mod.init_db()

    class Env:
        pass

    e = Env()
    e.session = session_mod
    e.models = models_mod
    e.CS = contractors_service_mod
    e.ES = export_service_mod
    e.coverage_service = coverage_service_mod
    e.overview_service = overview_service_mod
    e.contractors_route = contractors_route
    return e


@pytest.fixture()
def client(env):
    from fastapi import FastAPI
    from fastapi.testclient import TestClient

    app = FastAPI()
    app.include_router(env.contractors_route.router, prefix='/contractors')
    with TestClient(app) as c:
        yield c


def test_export_has_analysis_sheet_first_then_raw(client):
    """التصدير كان يعيد ٥٠٠ بسبب _filters_label/ES غير المعرّفتين — الآن يعمل
    وورقة التحليل هي الأولى، مبنية من نفس rows المصفّاة."""
    client.post('/contractors', json=dict(code='C1', name='مقاول واحد', projects=['الرشين']))
    client.post('/contractors/C1/entries', json=dict(
        date='2025-01-01', debit=0, credit=5000, description='مستخلص'))

    r = client.get('/contractors/export.xlsx')
    assert r.status_code == 200
    wb = load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames[0] == 'تحليل المقاولين'
    assert 'المقاولون' in wb.sheetnames

    ws = wb['تحليل المقاولين']
    text = '\n'.join(str(c.value) for row in ws.iter_rows() for c in row if c.value is not None)
    assert 'الإجماليات بالاتجاه' in text
    assert 'التوزيع بالمشروع' in text
    assert 'التوزيع بالحالة' in text
    assert 'أعلى ١٠' in text
    assert 'التصفية المطبَّقة' in text


def test_export_analysis_reflects_filtered_set_only(client):
    client.post('/contractors', json=dict(code='C1', name='مقاول أول', projects=['الرشين']))
    client.post('/contractors', json=dict(code='C2', name='مقاول ثاني', projects=['القصر']))
    client.post('/contractors/C1/entries', json=dict(
        date='2025-01-01', debit=0, credit=1000, description='مستخلص'))
    client.post('/contractors/C2/entries', json=dict(
        date='2025-01-01', debit=0, credit=9000, description='مستخلص'))

    r = client.get('/contractors/export.xlsx?project=الرشين')
    assert r.status_code == 200
    wb = load_workbook(io.BytesIO(r.content))
    ws = wb['تحليل المقاولين']
    text = '\n'.join(str(c.value) for row in ws.iter_rows() for c in row if c.value is not None)
    # المقاول C2 (٩٠٠٠) مستبعد بالتصفية — لا يظهر في «أعلى ١٠»
    assert 'مقاول ثاني' not in text
    assert 'الرشين' in text


def test_logo_degrades_gracefully_without_pillow(env, monkeypatch):
    """الشعار زينة — فشل تحميله (Pillow غائبة مثلاً) لا يجوز أن يُسقط التصدير."""
    monkeypatch.setattr(env.ES, '_logo_image', lambda: None)
    data = dict(rows=[], totals=dict(count=0, owedToContractors=0, owedToUs=0,
                                     balance=0, retentionHeld=0, byStatus={}))
    db = env.session.SessionLocal()
    try:
        out = env.ES.build_contractors_export_workbook(data, 'بلا تصفية', db)
    finally:
        db.close()
    wb = load_workbook(io.BytesIO(out))
    assert 'تحليل المقاولين' in wb.sheetnames


def test_totals_expose_derived_vs_reported_count(client, env):
    """م-٢٤: totals يفصح عن عدد المقاولين الذين اشتُقّ منهم owedToContractors فعلاً،
    بلا دمج الرقمين — انظر تعليق reported_balance في models.py."""
    client.post('/contractors', json=dict(code='C1', name='بحركات'))
    client.post('/contractors/C1/entries', json=dict(
        date='2025-01-01', debit=0, credit=1000, description='مستخلص'))
    client.post('/contractors', json=dict(code='C2', name='بلا حركات'))

    r = client.get('/contractors')
    body = r.json()
    assert body['totals']['derivedFromEntriesCount'] == 1
    assert body['totals']['count'] == 2


def test_contractor_coverage_reports_without_statement(env):
    """م-١٩: تغطية المقاولين — لم تكن موجودة إطلاقاً (grep contractor كان صفراً)."""
    db = env.session.SessionLocal()
    try:
        c1 = env.models.Contractor(code='C1', name='له حركات')
        c2 = env.models.Contractor(code='C2', name='بلا حركات', reported_balance=-5000)
        db.add_all([c1, c2])
        db.flush()
        db.add(env.models.ContractorEntry(
            contractor_id=c1.id, date=__import__('datetime').date(2025, 1, 1),
            debit=0, credit=100, description='x', kind='claim', source='manual'))
        db.commit()

        import datetime as _dt
        cov = env.coverage_service.contractor_coverage(db, today=_dt.date(2025, 1, 10))
        assert cov['totals']['contractors'] == 2
        assert cov['totals']['withoutData'] == 1
        states = {r['code']: r['state'] for r in cov['rows']}
        assert states['C1'] == 'ok'
        assert states['C2'] == 'none'
    finally:
        db.close()


def test_status_filter_and_by_status_totals(client):
    """م-٢٥ (واجهة الحالة): status/statusNote في الصف · تصفية status · byStatus."""
    client.post('/contractors', json=dict(code='C1', name='نشط'))
    client.post('/contractors', json=dict(code='C2', name='متوقف'))
    client.post('/contractors/C1/entries', json=dict(
        date='2025-01-01', debit=0, credit=1000, description='مستخلص'))
    client.post('/contractors/C2/entries', json=dict(
        date='2025-01-01', debit=0, credit=2000, description='مستخلص'))

    r = client.put('/contractors/C2', json=dict(status='paused', statusNote='بانتظار قرار'))
    assert r.status_code == 200
    assert r.json()['status'] == 'paused'
    assert r.json()['statusNote'] == 'بانتظار قرار'

    r = client.get('/contractors?status=paused')
    body = r.json()
    assert body['count'] == 1
    assert body['rows'][0]['code'] == 'C2'
    assert body['totals']['byStatus']['paused']['count'] == 1
    assert body['totals']['byStatus']['paused']['owedToContractors'] == 2000.0


def test_status_rejects_unknown_value_with_arabic_422(client):
    client.post('/contractors', json=dict(code='C1', name='مقاول'))
    r = client.put('/contractors/C1', json=dict(status='not_a_real_status'))
    assert r.status_code == 422
    assert 'حالة غير صالحة' in r.json()['detail']


def test_status_query_filter_rejects_unknown_value(client):
    r = client.get('/contractors?status=bogus')
    assert r.status_code == 422


def test_receivables_total_collected_uses_decimal_not_float_drift(tmp_path, monkeypatch):
    """م-١١: كان sum() خاماً على float بينما كل مسار تحصيلات آخر يلفّ الصف بـ D() —
    ثلاث قيم تُسبّب انحرافاً معروفاً بجمع float مباشر (0.1+0.2+0.3 != 0.6)."""
    monkeypatch.setenv('EGCO_DATA_DIR', str(tmp_path / 'data'))
    import app.core.config as config_mod
    importlib.reload(config_mod)
    import app.db.session as session_mod
    importlib.reload(session_mod)
    import app.services.receivables_service as receivables_service_mod
    importlib.reload(receivables_service_mod)
    session_mod.init_db()

    class _Row:
        def __init__(self, unit, client, amount, status):
            self.unit, self.client, self.amount, self.status = unit, client, amount, status
            self.project, self.due_date, self.collected_on = '', None, None

    def _fake_parser(path):
        rows = [_Row('U1', 'ع1', 0.1, 'collected'),
               _Row('U2', 'ع2', 0.2, 'collected'),
               _Row('U3', 'ع3', 0.3, 'collected')]
        return dict(receivables=rows, issues=[])

    monkeypatch.setitem(receivables_service_mod._PARSERS, 'fake_source', _fake_parser)
    monkeypatch.setattr('app.services.import_service.backup_db', lambda: None, raising=False)

    db = session_mod.SessionLocal()
    try:
        result = receivables_service_mod.import_receivables(db, 'fake/path', source='fake_source')
        assert result['totalCollected'] == 0.6
    finally:
        db.close()


def test_overview_alerts_on_contractors_without_statement(env):
    """م-١٩: أكبر فجوة (مقاولون بلا كشف حركات) يجب أن تظهر تنبيهاً صريحاً بالحجم،
    لا أن تُصمت عنها اللوحة بينما تُفصح عن فجوة الموردين الأصغر."""
    db = env.session.SessionLocal()
    try:
        c = env.models.Contractor(code='C1', name='بلا كشف', reported_balance=-750000)
        db.add(c)
        db.commit()

        ov = env.overview_service.overview(db)
        assert ov['contractors']['withoutStatementCount'] == 1
        assert ov['contractors']['withoutStatementOwed'] == 750000.0
        kinds = [a.get('kind') for a in ov['alerts']]
        assert 'contractors_without_statement' in kinds
    finally:
        db.close()
